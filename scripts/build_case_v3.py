#!/usr/bin/env python3
"""Reproducible, docs-only Intelligence Cup 2026 quality-upgrade submission exporter.
Reads the GitHub canonical main/Q5 Markdown and scenario JSON; does NOT run a Qwen model.
"""
import hashlib, json, re, shutil, subprocess, sys, os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor as PC
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(sys.argv[1]).resolve() if len(sys.argv)>1 else Path(__file__).resolve().parents[1]
OUT = (Path(sys.argv[2]) if len(sys.argv)>2 else ROOT / 'dist').resolve()
OUT.mkdir(parents=True,exist_ok=True)
SRC=ROOT/'submissions'/'main-solution.ru.md'; Q5=ROOT/'submissions'/'05-effect-metrics.ru.md'
SOURCES=ROOT/'docs'/'business'/'scenario-inputs.json'
main=SRC.read_text('utf-8'); q5=Q5.read_text('utf-8'); finance=json.loads(SOURCES.read_text('utf-8'))
assert 5000<len(main)<=15000, f'Main character limit: {len(main)}'
assert all(s in main for s in ('1 532 277','Wilson','BM25','6,35%')), 'Upgrade incomplete'
assert all(s in q5 for s in ('Wilson','246','BM25')), 'Q5 upgrade incomplete'

NAVY=PC(16,49,76); TEAL=PC(0,145,143); ORANGE=PC(245,106,43)
BG=PC(249,251,253); PALE=PC(230,242,250); LAV=PC(239,234,252); MINT=PC(225,246,242); PEACH=PC(255,238,225)
GRAY=PC(87,110,128); WHITE=PC(255,255,255); LINE=PC(215,228,236)

# DOCX: native editable headings/tables; standalone answer (Q1-4,6 only)
doc=Document();sec=doc.sections[0]
sec.page_height=Inches(11.69);sec.page_width=Inches(8.27)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7)
sec.left_margin=Inches(.71);sec.right_margin=Inches(.71)
styles=doc.styles;normal=styles['Normal'];normal.font.name='DejaVu Sans';normal.font.size=Pt(9.65);normal.font.color.rgb=RGBColor(28,52,74)
normal.paragraph_format.space_after=Pt(4);normal.paragraph_format.line_spacing=1.11
for name,size,bold,color in [('Title',17,True,(16,49,76)),('Heading 1',13,True,(16,49,76)),('Heading 2',11,True,(0,127,132)),('Heading 3',10,True,(0,127,132))]:
 st=styles[name];st.font.name='DejaVu Sans';st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(*color)
 st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(4)

header=sec.header.paragraphs[0];header.text='INTELLIGENCE CUP 2026  •  GAZPROM NEFT  /  PROJECT CONCEPT'
header.alignment=WD_PARAGRAPH_ALIGNMENT.RIGHT
for r in header.runs:r.font.size=Pt(7);r.font.color.rgb=RGBColor(94,119,135)
foot=sec.footer.paragraphs[0];foot.text='ПРОЕКТНАЯ КОНЦЕПЦИЯ  •  СЦЕНАРИИ ≠ ИЗМЕРЕННЫЕ РЕЗУЛЬТАТЫ'
for r in foot.runs:r.font.size=Pt(7);r.font.color.rgb=RGBColor(110,128,143)

def runs(p,s):
  # Inline strong and backtick; no raw markup in the visible DOCX text.
  segs=re.split(r'(\*\*[^*]+\*\*|`[^`]+`)',s)
  for z in segs:
    if not z:continue
    if z.startswith('**') and z.endswith('**'): rr=p.add_run(z[2:-2]);rr.bold=True
    elif z.startswith('`') and z.endswith('`'):rr=p.add_run(z[1:-1]);rr.font.name='DejaVu Sans Mono';rr.font.size=Pt(8)
    else:rr=p.add_run(z)

def table_from_rows(rows):
  cols=max(len(r) for r in rows)
  t=doc.add_table(rows=0,cols=cols);t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
  for ix,parts in enumerate(rows):
    cells=t.add_row().cells
    for j,val in enumerate(parts):
      cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
      p=cells[j].paragraphs[0];p.paragraph_format.space_after=Pt(1)
      runs(p,val)
      for rr in p.runs:rr.font.size=Pt(7.6)
      if ix==0:
        cells[j].fill if hasattr(cells[j],'fill') else None
        tcPr=cells[j]._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'174767');tcPr.append(sh)
        for rr in p.runs:rr.font.color.rgb=RGBColor(255,255,255);rr.bold=True
    if ix==0:
      trPr=t.rows[0]._tr.get_or_add_trPr();tblH=OxmlElement('w:tblHeader');tblH.set(qn('w:val'),'true');trPr.append(tblH)
  doc.add_paragraph().paragraph_format.space_after=Pt(1)

lines=main.splitlines();in_code=False;table_rows=[]
for ln in lines:
  if ln.startswith('```'):
    if table_rows:table_from_rows(table_rows);table_rows=[]
    in_code=not in_code;continue
  if ln.startswith('|') and ln.endswith('|'):
    cols=[c.strip() for c in ln[1:-1].split('|')]
    if cols and all(re.fullmatch(r':?-{2,}:?',c or '') for c in cols):continue
    table_rows.append(cols);continue
  if table_rows:table_from_rows(table_rows);table_rows=[]
  if not ln.strip():continue
  if in_code:
    p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.17);p.paragraph_format.space_after=Pt(0)
    rr=p.add_run(ln);rr.font.name='DejaVu Sans Mono';rr.font.size=Pt(7.6);rr.font.color.rgb=RGBColor(34,77,96)
  elif ln.startswith('# '):p=doc.add_paragraph(style='Title');runs(p,ln[2:])
  elif ln.startswith('## '):p=doc.add_paragraph(style='Heading 1');runs(p,ln[3:])
  elif ln.startswith('### '):p=doc.add_paragraph(style='Heading 2');runs(p,ln[4:])
  else:
    bullet=re.match(r'^(?:\d+\.|[-*])\s+(.*)$',ln)
    p=doc.add_paragraph(style='Normal');
    if bullet:p.paragraph_format.left_indent=Inches(.14);runs(p,'• '+bullet.group(1))
    else:runs(p,ln)
if table_rows:table_from_rows(table_rows)
docx=OUT/'Intelligence_Cup_2026_Main_Q1-4_Q6_V3.docx';doc.save(docx)
shutil.copyfile(Q5,OUT/'Intelligence_Cup_2026_Q5_Effect_Metrics_V3.md')

# PPTX: original-inspired light, editable native vectors; 10 slides exactly.
prs=Presentation();prs.slide_width=PI(13.333);prs.slide_height=PI(7.5)

def shape(sl,x,y,w,h,fill=WHITE,stroke=LINE,r=.12):
  sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,PI(x),PI(y),PI(w),PI(h))
  sh.fill.solid();sh.fill.fore_color.rgb=fill
  sh.line.fill.background() if stroke is None else None
  if stroke is not None:sh.line.color.rgb=stroke;sh.line.width=PP(.75)
  sh.adjustments[0]=r
  return sh

def txt(sl,s,x,y,w,h,size=17,color=NAVY,bold=False,align=PP_ALIGN.LEFT,margin=.02):
  box=sl.shapes.add_textbox(PI(x),PI(y),PI(w),PI(h));tf=box.text_frame;tf.clear();tf.word_wrap=True
  tf.margin_left=tf.margin_right=PI(margin);tf.margin_top=tf.margin_bottom=PI(margin)
  for j,line in enumerate(s.split('\n')):
    p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.alignment=align;p.space_after=PP(4)
    p.font.name='DejaVu Sans';p.font.size=PP(size);p.font.bold=bold;p.font.color.rgb=color
  return box

def base(title,kicker,n,desc=None):
  sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=BG
  shape(sl,0,0,13.333,.09,ORANGE,None,0)
  txt(sl,kicker.upper(),.48,.31,11,.35,10,TEAL,True)
  txt(sl,f'{n:02d} / 10',11.93,.31,.92,.36,10,GRAY,False,PP_ALIGN.RIGHT)
  txt(sl,title,.5,.82,12.1,.85,26,NAVY,True)
  if desc:txt(sl,desc,.53,1.76,12,.59,12,GRAY)
  txt(sl,'INTELLIGENCE CUP 2026  •  ЦИФРОВЫЕ ПРОДУКТЫ  •  КОНЦЕПЦИЯ',.48,7.17,8,.16,7.2,GRAY)
  return sl

def card(sl,x,y,w,h,head,body,accent=TEAL,fill=PALE,headsize=14,bodysize=14):
  shape(sl,x,y,w,h,fill,LINE);shape(sl,x+.16,y+.16,w-.32,.47,WHITE,None)
  txt(sl,head,x+.27,y+.25,w-.54,.29,headsize,accent,True)
  txt(sl,body,x+.24,y+.89,w-.48,h-1.02,bodysize,NAVY)

sl=base('От обращений пользователей к развитию цифровых продуктов','задача конкурса',1,'Единый подход для десятков внутренних и внешних цифровых продуктов. Вебнефть — лишь пример из PDF.')
shape(sl,.55,2.63,12.2,2.72,PALE,LINE);txt(sl,'PRODUCT FEEDBACK INTELLIGENCE',.95,3.11,11.2,.48,25,TEAL,True,PP_ALIGN.CENTER)
txt(sl,'Тикеты + звонки → доказательства → полезный сигнал → product team → feedback',.92,3.95,11.4,.81,18,NAVY,True,PP_ALIGN.CENTER)
txt(sl,'HYBRID AI  •  INTERNAL QWEN  •  EVIDENCE FIRST  •  NOT DEPLOYED',.8,6.29,11.8,.33,12,GRAY,False,PP_ALIGN.CENTER)

sl=base('Поддержка уже работает. Но закономерности теряются.','проблема',2,'Не подменяем текущий процесс: его частота и трудозатраты в условиях не раскрыты.')
card(sl,.55,2.52,3.75,2.6,'ИСТОЧНИК','Заявки\nЗвонки и файлы\n1 год хранения',TEAL,PALE)
card(sl,4.8,2.52,3.75,2.6,'АГЕНТ','Темы + evidence\nБез добавления нагрузки\nБез внешних API',PC(109,75,174),LAV)
card(sl,9.05,2.52,3.75,2.6,'РЕЗУЛЬТАТ','Нужная product team\nРешение + feedback\nИзмеримый срок',TEAL,MINT)
txt(sl,'≈40 менеджеров — весь контур; число сотрудников ≠ число тикетов',1.3,6.17,10.9,.45,14,NAVY,True,PP_ALIGN.CENTER)

sl=base('Webneft — 7 обращений, но не 7 инцидентов','бизнес-пример',3,'Семь текстовых примеров из PDF позволяют создать taxonomy, но не доказать повторяемость.')
items=[('01–02','Договор / оферта','ИП вместо ООО; кнопка «Далее»'),('03','Транспорт','Машина исчезла / «уже существует»'),('04','Отгрузка','Срочная смена номера → действующий SLA'),('05','Электронная очередь','HTTP 404 / не найден слот'),('06–07','Доступ','Старый договор / восстановление SMS')]
for i,(no,name,desc) in enumerate(items):
 y=2.48+i*.78;shape(sl,.6,y,12.05,.65,PALE if i%2==0 else WHITE,LINE)
 txt(sl,no,.83,y+.15,1.2,.3,13,TEAL,True);txt(sl,name,2.1,y+.14,3.1,.33,13,NAVY,True);txt(sl,desc,5.2,y+.14,7.05,.37,12,NAVY)

sl=base('5 логических функций — 3 процесса MVP','архитектура / нагрузка',4,'Каждая величина ниже — сценарная, не реальные данные компании и не число GPU.')
heads=['INGESTION','ASR / VL','QWEN + BM25','AGGREGATE','DELIVERY']; nums=['300','4 + 1','2','0,25','7,5'];subs=['тикетов / ч','условных слота','text slots','DB writes / s','дайджестов / ч']
for i in range(5):
 x=.55+i*2.55;card(sl,x,2.5,2.4,2.75,heads[i],nums[i]+'\n'+subs[i],[TEAL,PC(116,75,178),TEAL,PC(78,109,175),ORANGE][i],[PALE,LAV,MINT,PALE,PEACH][i],11,18)
shape(sl,.6,5.72,12.07,.83,WHITE,LINE)
txt(sl,'≈40 × 20 обращений/день ÷ 8 часов × пик 3 = 300/ч. RTF/VRAM/latency измерять отдельно.',.92,6.03,11.45,.35,12,NAVY,True,PP_ALIGN.CENTER)

sl=base('Данные остаются в закрытом корпоративном контуре','security / trust',5,'Аудио и изображения могут содержать персональные данные до извлечения текста.')
for i,(h,b) in enumerate([('TICKET DB','Звонки\nФайлы'),('QWEN-ASR / VL','Только доступный\nmedia asset'),('MINIMIZATION','PII / secrets\nTrusted aliases'),('INTERNAL QWEN','Topic / intent\nГипотезы')]):
 card(sl,.55+i*3.18,2.5,3.02,2.93,h,b,TEAL if i!=1 else PC(106,71,178),[PALE,LAV,PEACH,MINT][i],13,16)
shape(sl,1.1,5.83,11.13,.66,WHITE,LINE)
txt(sl,'Backend: product_id + taxonomy_version + RBAC + team registry  ≠  LLM',1.4,6.05,10.6,.28,13,NAVY,True,PP_ALIGN.CENTER)

sl=base('LLM предлагает — доверенные tools проверяют','tools / permissions',6,'Ни SQL/HTTP от модели, ни внешних API, ни права на изменения заказов и аккаунтов.')
card(sl,.55,2.43,3.77,3.52,'READ / INGEST','list_new_tickets\nget_ticket_context\ntranscribe_recording\nanalyze_attachment',TEAL,PALE,13,14)
card(sl,4.77,2.43,3.77,3.52,'ANALYZE / VERIFY','find_related_requests\naggregate_recurrence\nresolve_product_team',PC(114,75,178),LAV,13,14)
card(sl,8.99,2.43,3.77,3.52,'WRITE — POLICY ONLY','create_insight_draft\nsend_internal_digest\nrecord_product_feedback',TEAL,MINT,13,14)
txt(sl,'Unknown team → human triage  •  Unknown coverage → observed count, не «рост»',.8,6.49,11.9,.35,12,ORANGE,True,PP_ALIGN.CENTER)

sl=base('Один проверяемый сигнал для product manager','синтетический e2e',7,'Три ID полностью выдуманы; исходный PDF содержит один пример близкой темы.')
card(sl,.5,2.52,3.77,3.54,'SOURCE • SYNTHETIC','SYN-QUEUE-001: 404\nSYN-QUEUE-002: «нет времени»\nSYN-QUEUE-003: нет слота',TEAL,PALE,12,14)
card(sl,4.75,2.52,3.78,3.54,'MODEL → BACKEND','Qwen: тема (гипотеза)\nBackend: observed count=3\nCoverage: unknown',PC(112,77,171),LAV,12,14)
card(sl,9,2.52,3.78,3.54,'PRODUCT INSIGHT • DRAFT','Evidence: 3 SYN IDs\nRoot cause: unknown\nDelivery: not sent\nFeedback: pending',TEAL,MINT,12,14)
shape(sl,.75,6.27,11.86,.53,PEACH,None);txt(sl,'Ни общая причина, ни тренд НЕ доказаны.',1.0,6.42,11.28,.3,13,ORANGE,True,PP_ALIGN.CENTER)

sl=base('Почему Hybrid C? Сначала проверим, лучше ли он BM25.','выбор архитектуры / эксперимент',8,'Одинаковый temporal/product holdout, независимые метки, без предположения о выигрыше модели.')
for i,(h,b) in enumerate([('A • MANUAL','Текущий процесс\nЧасы и ручной baseline'),('B • RULES + BM25','Лексические темы\nДешевый и надежный fallback'),('C • HYBRID QWEN','Перефразы и звонки\nТолько при измеренном приросте')]):
 card(sl,.57+i*4.27,2.4,4.05,2.68,h,b,[PC(82,109,140),TEAL,PC(111,77,167)][i],[PALE,MINT,LAV][i],13,14)
shape(sl,.69,5.49,11.95,.98,WHITE,LINE)
txt(sl,'Проверка: gold TP/FP, ложные объединения, net reviewer minutes, coverage, SLA',.97,5.77,11.4,.37,12,NAVY,True,PP_ALIGN.CENTER)
txt(sl,'Precision ≥0,80 — цель; GO лишь по нижней 95% Wilson CI ≥0,80 после согласования.',.8,6.58,11.9,.36,10,ORANGE,True,PP_ALIGN.CENTER)

sl=base('5 минут после qualification ≠ 96 часов от первого ticket','метрики',9,'Оба времени — иллюстрация, а не измеренный результат компании.')
for i,(date,h,field) in enumerate([('26.10 09:00','Первое наблюдение','first_relevant_ingested_at'),('30.10 09:00','Evidence достаточно','qualified_at'),('30.10 09:05','Доставка команде','sent_at')]):
 x=.6+i*4.26;shape(sl,x,2.47,3.91,1.38,[PALE,LAV,MINT][i],LINE)
 txt(sl,date,x+.2,2.75,3.53,.37,17,NAVY,True,PP_ALIGN.CENTER);txt(sl,h,x+.2,3.34,3.51,.32,11,GRAY,False,PP_ALIGN.CENTER)
 txt(sl,field,x,4.0,3.91,.38,10,GRAY,False,PP_ALIGN.CENTER)
shape(sl,.65,4.62,5.84,.81,PALE,None);shape(sl,6.83,4.62,5.84,.81,MINT,None)
txt(sl,'E2E = 96 ч 05 мин',.9,4.9,5.27,.31,17,NAVY,True,PP_ALIGN.CENTER)
txt(sl,'TPI = 5 мин',7.11,4.9,5.28,.31,17,TEAL,True,PP_ALIGN.CENTER)
txt(sl,'Pilot goal p90 TPI ≤4ч. Undelivered/eligible, backlog и missed clocks — показывать отдельно.',.8,6.02,11.8,.46,11,ORANGE,True,PP_ALIGN.CENTER)

sl=base('3 года: сценарная ценность, не прибыль и не ROI','финансовая модель / macro',10,'Все PV — учебные сценарии высвобожденного времени; затрат C и эффекта Qwen не измеряли.')
for i,(name,amount,color) in enumerate([('Консервативный','123 тыс. ₽',PC(78,113,159)),('Базовый','1,64 млн ₽',TEAL),('Оптимистичный','7,39 млн ₽',ORANGE)]):
 y=2.42+i*.86;shape(sl,.65,y,12.05,.67,WHITE,LINE)
 txt(sl,name,.9,y+.19,3.2,.29,14,NAVY,True);shape(sl,4.05,y+.21,2.0+i*2.11,.22,color,None)
 txt(sl,amount,9.83,y+.18,2.53,.32,15,NAVY,True,PP_ALIGN.RIGHT)
shape(sl,.65,5.17,12.05,.76,PALE,LINE)
txt(sl,'Россия на 30.10.2026: ставка ЦБ 14%; инфляция сентябрь 6,35% г/г.',.9,5.38,11.57,.33,12,NAVY,True,PP_ALIGN.CENTER)
shape(sl,.65,6.05,12.05,.73,PEACH,None)
txt(sl,'Стресс r=14%: PV 1,53 млн ₽  •  r ≠ WACC  •  Hybrid vs BM25 — только пилот',.93,6.26,11.51,.3,11,ORANGE,True,PP_ALIGN.CENTER)
txt(sl,'Источники: Банк России (30.10.2026); Росстат (ИПЦ сентябрь 2026)',.95,6.86,11.4,.20,8,GRAY,False,PP_ALIGN.CENTER)

slides=OUT/'Intelligence_Cup_2026_Presentation_V3.pptx';prs.save(slides)
# Pandoc not needed: original Markdown is the source of DOCX. Export PDFs via LO.
lo=shutil.which('libreoffice') or shutil.which('soffice')
if lo:
 for p in (docx,slides):
  cmd=[lo,'-env:UserInstallation=file:///tmp/lo_icup_v3_'+str(os.getpid()),'--headless','--convert-to','pdf','--outdir',str(OUT),str(p)]
  r=subprocess.run(cmd,capture_output=True,text=True,timeout=150)
  if r.returncode!=0:raise RuntimeError(str(cmd)+'\n'+r.stdout+'\n'+r.stderr)
else:
 print('WARNING: LibreOffice not installed; DOCX/PPTX produced, PDF conversion skipped')

manifest={'main_markdown_characters':len(main),'source_note':'Official case p5 max 15,000 Unicode chars, item5 separate .md', 'files':{}}
for p in sorted(OUT.iterdir()):
 if p.is_file() and p.name!='manifest.json':manifest['files'][p.name]={'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
(OUT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n','utf-8')
print(json.dumps({'main_chars':len(main),'slides':len(prs.slides),'exports':list(manifest['files'])},ensure_ascii=False))